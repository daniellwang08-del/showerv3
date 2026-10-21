"""Studio features: heading styles, renamed sections, extra sections, date formats, themes."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from pydantic import ValidationError

from app.models.database import User
from app.models.resume_design_schemas import ExtraItem, ExtraSection, ResumeDesign, section_title
from app.services.resume_custom_theme_service import snapshot_design_as_theme
from app.services.resume_design_compiler import _heading, _render_extra_section, compile_design
from app.services.resume_themes import THEME_PRESETS, build_catalog
from app.utils.flexible_date import date_style, format_flexible_date, format_flexible_period


def _user() -> User:
    return User(
        id="u1",
        email="a@x.com",
        name_first="Ada",
        name_last="Lovelace",
        profile_title="Engineer",
        profile_email="ada@example.com",
        profile_summary="Builds things.",
        technical_skills=[{"category": "Languages", "skills": "Python"}],
        work_experience=[
            {"company_name": "Acme", "job_title": "Engineer", "period_start": "2021-03", "period_end": ""},
        ],
        education=[],
        certificates=[],
        extra=[],
    )


def _projects(**kw) -> ExtraSection:
    return ExtraSection(
        id="x-proj01",
        kind="projects",
        items=[
            ExtraItem(title="Compiler", subtitle="Rust", date_start="2022-01", date_end="2022-06", bullets=["Wrote it."]),
            ExtraItem(title="Notebook", url="example.com/nb"),
        ],
        **kw,
    )


@pytest.mark.parametrize(
    "fmt,expected",
    [
        ("mon_yyyy", "Mar 2021 - Now"),
        ("month_yyyy", "March 2021 - Now"),
        ("mm_yyyy", "03/2021 - Now"),
        ("yyyy", "2021 - Now"),
    ],
)
def test_date_formats_and_present_label(fmt: str, expected: str) -> None:
    with date_style(fmt, "Now"):
        assert format_flexible_period("2021-03", "") == expected
        assert format_flexible_date("present") == "Now"
    assert format_flexible_period("2021-03", "") == "Mar 2021 - Present"


def test_year_only_collapses_identical_years() -> None:
    with date_style("yyyy"):
        assert format_flexible_period("2021-01", "2021-09") == "2021"
    assert format_flexible_period("2021", "2021") == "2021 - 2021"


def test_titles_are_cleaned_and_resolved() -> None:
    design = ResumeDesign.model_validate(
        {"sections": {"titles": {"experience": "  Professional   Experience ", "skills": "   "}}}
    )
    assert design.sections.titles == {"experience": "Professional Experience"}
    design.extra_sections = [_projects()]
    assert section_title(design, "experience") == "Professional Experience"
    assert section_title(design, "education") == "Education"
    assert section_title(design, "x-proj01") == "Projects"


def test_extra_section_ids_are_validated() -> None:
    with pytest.raises(ValidationError):
        ExtraSection(id="projects")


@pytest.mark.parametrize(
    "style,sides,shaded,white",
    [
        ("plain", set(), False, False),
        ("underline", {"bottom"}, False, False),
        ("overline", {"top"}, False, False),
        ("double", {"top", "bottom"}, False, False),
        ("left_bar", {"left"}, False, False),
        ("tint", {"top", "left", "bottom", "right"}, True, False),
        ("solid", {"top", "left", "bottom", "right"}, True, True),
        ("boxed", {"top", "left", "bottom", "right"}, False, False),
    ],
)
def test_heading_styles(style: str, sides: set[str], shaded: bool, white: bool) -> None:
    design = ResumeDesign()
    design.layout.heading_style = style  # type: ignore[assignment]
    design.typography.heading_font_family = "Lora"
    doc = Document()
    _heading(doc, "Experience", design)
    p_pr = doc.paragraphs[0]._p.pPr
    borders = p_pr.find(qn("w:pBdr"))
    found = {el.tag.split("}")[1] for el in borders} if borders is not None else set()
    assert found == sides
    assert (p_pr.find(qn("w:shd")) is not None) == shaded
    run = doc.paragraphs[0].runs[0]
    assert run.font.name == "Lora"
    assert (str(run.font.color.rgb) == "FFFFFF") == white


def test_legacy_accent_rule_still_underlines() -> None:
    design = ResumeDesign()
    assert design.layout.effective_heading_style == "underline"
    design.layout.accent_rule = False
    assert design.layout.effective_heading_style == "plain"


@pytest.mark.parametrize("layout", ["entries", "list", "inline", "columns"])
def test_extra_section_layouts_render_every_item(layout: str) -> None:
    design = ResumeDesign()
    section = _projects(layout=layout)
    doc = Document()
    _render_extra_section(doc, design, section)
    xml = doc.element.xml
    assert "Projects" in doc.paragraphs[0].text.title()
    assert "Compiler" in xml and "Notebook" in xml
    if layout == "columns":
        assert len(doc.tables) == 1
    if layout == "entries":
        assert "Wrote it." in xml
        assert "Jan 2022 - Jun 2022" in xml


def test_compile_uses_titles_extras_and_date_format(tmp_path: Path) -> None:
    design = ResumeDesign()
    design.sections.titles = {"experience": "Professional Experience"}
    design.sections.date_format = "month_yyyy"
    design.extra_sections = [_projects(), ExtraSection(id="x-empty1", kind="awards")]
    design.layout.section_order = [*design.layout.section_order, "x-proj01", "x-empty1"]
    out = tmp_path / "r.docx"
    compile_design(design, _user(), out)
    xml = zipfile.ZipFile(out).read("word/document.xml").decode()
    assert "PROFESSIONAL EXPERIENCE" in xml
    assert "PROJECTS" in xml and "Compiler" in xml
    assert "January 2022 - June 2022" in xml
    assert "AWARDS" not in xml


def test_theme_snapshot_drops_personal_sections() -> None:
    design = ResumeDesign()
    design.extra_sections = [_projects()]
    design.sections.titles = {"experience": "Work"}
    design.layout.section_order = [*design.layout.section_order, "x-proj01"]
    snap = snapshot_design_as_theme(design, theme_id="custom-1")
    assert snap.extra_sections == []
    assert snap.sections.titles == {}
    assert "x-proj01" not in snap.layout.section_order


def test_catalog_themes_have_categories_and_new_looks() -> None:
    catalog = build_catalog()
    ids = {t.id for t in catalog.themes}
    assert {"harvard", "corporate", "consultant", "engineer", "impact", "scholar"} <= ids
    assert {t.category for t in THEME_PRESETS} == {"classic", "modern", "technical", "executive", "minimal"}
