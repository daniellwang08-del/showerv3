"""Certificate issued_at + url fields round-trip through DOCX rendering."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

from app.models.resume_design_schemas import ResumeDesign
from app.services.resume_design_compiler import _certificate_rows, _render_certificates


class _FakeUser:
    certificates = [
        {"name": "AWS Certified Solutions Architect", "issued_at": "2024-06", "url": "https://credly.com/badges/abc"},
        {"name": "DevOps Engineer", "issued_at": "", "url": ""},
    ]


def test_certificate_rows_include_optional_fields():
    rows = _certificate_rows(_FakeUser())  # type: ignore[arg-type]
    assert rows[0]["issued_at"] == "2024-06"
    assert rows[0]["url"] == "https://credly.com/badges/abc"
    assert rows[1]["name"] == "DevOps Engineer"


def test_render_certificates_hyperlinks_name_when_url_present():
    design = ResumeDesign()
    doc = Document()
    _render_certificates(
        doc,
        design,
        [
            {
                "name": "AWS Certified Solutions Architect",
                "issued_at": "2024-06",
                "url": "https://credly.com/badges/abc",
            }
        ],
    )
    targets = []
    texts = []
    for p in doc.paragraphs:
        texts.append(p.text)
        for hl in p._p.iter(qn("w:hyperlink")):
            r_id = hl.get(qn("r:id"))
            if r_id:
                targets.append(p.part.rels[r_id].target_ref)
    assert any("credly.com/badges/abc" in t for t in targets)
    assert any("2024-06" in t for t in texts)

    buf = BytesIO()
    doc.save(buf)
    reloaded = Document(BytesIO(buf.getvalue()))
    assert any(
        "credly.com/badges/abc" in p.part.rels[hl.get(qn("r:id"))].target_ref
        for p in reloaded.paragraphs
        for hl in p._p.iter(qn("w:hyperlink"))
        if hl.get(qn("r:id"))
    )
