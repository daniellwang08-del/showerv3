"""Regression: LinkedIn/GitHub hyperlink targets must survive DOCX/PDF text extraction."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.models.profile_schemas import ResumeExtractedDraft
from app.services.resume_parse_service import (
    _fill_missing_contact_from_text,
    docx_to_plain_text,
    pdf_to_plain_text_fitz,
    pymupdf_available,
)


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _docx_with_label_hyperlink(label: str, href: str) -> bytes:
    doc = Document()
    p = doc.add_paragraph("Jane Doe\n")
    p2 = doc.add_paragraph("Contact: ")
    _add_hyperlink(p2, href, label)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_extracts_linkedin_href_when_visible_text_is_only_label():
    raw = _docx_with_label_hyperlink("LinkedIn", "https://www.linkedin.com/in/jane-doe-123")
    text = docx_to_plain_text(raw)
    assert "LinkedIn" in text
    assert "linkedin.com/in/jane-doe-123" in text.lower()

    draft = ResumeExtractedDraft()
    notes = _fill_missing_contact_from_text(draft, text)
    assert draft.linkedin_url == "https://www.linkedin.com/in/jane-doe-123"
    assert any("LinkedIn" in n for n in notes)


def test_docx_extracts_linkedin_from_page_header():
    """Contact icon rows often live in the header — body-only extraction misses them."""
    doc = Document()
    section = doc.sections[0]
    hp = section.header.paragraphs[0]
    hp.text = "Jane Doe  "
    _add_hyperlink(hp, "https://www.linkedin.com/in/header-jane", "LinkedIn")
    doc.add_paragraph("Senior Engineer")
    doc.add_paragraph("Work experience only in the body")
    buf = BytesIO()
    doc.save(buf)
    text = docx_to_plain_text(buf.getvalue())
    assert "linkedin.com/in/header-jane" in text.lower()

    draft = ResumeExtractedDraft()
    notes = _fill_missing_contact_from_text(draft, text)
    assert draft.linkedin_url == "https://www.linkedin.com/in/header-jane"
    assert any("LinkedIn" in n for n in notes)


def test_docx_extracts_github_href_when_visible_text_is_only_label():
    raw = _docx_with_label_hyperlink("GitHub", "https://github.com/jane-doe")
    text = docx_to_plain_text(raw)
    assert "github.com/jane-doe" in text.lower()

    draft = ResumeExtractedDraft()
    _fill_missing_contact_from_text(draft, text)
    assert draft.github_url == "https://github.com/jane-doe"


def test_docx_does_not_duplicate_url_already_visible_as_text():
    raw = _docx_with_label_hyperlink(
        "linkedin.com/in/jane-doe-123",
        "https://www.linkedin.com/in/jane-doe-123",
    )
    text = docx_to_plain_text(raw)
    assert text.lower().count("linkedin.com/in/jane-doe-123") == 1


@pytest.mark.skipif(not pymupdf_available(), reason="PyMuPDF required")
def test_pdf_extracts_linkedin_uri_annotation_when_visible_text_is_only_label():
    import fitz

    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 72), "Jane Doe", fontsize=16)
    page.insert_text((72, 100), "jane@example.com", fontsize=11)
    page.insert_text((72, 120), "LinkedIn", fontsize=11)
    page.insert_link(
        {
            "kind": fitz.LINK_URI,
            "from": fitz.Rect(72, 108, 140, 128),
            "uri": "https://www.linkedin.com/in/jane-doe-123",
        }
    )
    raw = doc.tobytes()
    doc.close()

    text = pdf_to_plain_text_fitz(raw)
    assert "LinkedIn" in text
    assert "linkedin.com/in/jane-doe-123" in text.lower()

    draft = ResumeExtractedDraft()
    notes = _fill_missing_contact_from_text(draft, text)
    assert draft.linkedin_url == "https://www.linkedin.com/in/jane-doe-123"
    assert any("LinkedIn" in n for n in notes)


@pytest.mark.skipif(not pymupdf_available(), reason="PyMuPDF required")
def test_pdf_visible_url_text_still_recovers_without_annotation():
    import fitz

    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 72), "Jane Doe", fontsize=16)
    page.insert_text((72, 100), "linkedin.com/in/visible-handle", fontsize=11)
    raw = doc.tobytes()
    doc.close()

    text = pdf_to_plain_text_fitz(raw)
    draft = ResumeExtractedDraft()
    _fill_missing_contact_from_text(draft, text)
    assert draft.linkedin_url == "https://www.linkedin.com/in/visible-handle"
