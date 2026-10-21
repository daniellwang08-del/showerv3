"""Optional resume headshot: schema, DOCX placement and header band sizing."""

from __future__ import annotations

import base64
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image
from pydantic import ValidationError

from app.models.profile_schemas import ProfilePhotoRequest
from app.models.resume_design_schemas import PHOTO_SIZE_PT, ResumeDesign, ResumePhoto
from app.services.resume_design_compiler import _header_band_height_pt, _render_header


def _png_data_url() -> str:
    buf = BytesIO()
    Image.new("RGB", (40, 40), (120, 140, 160)).save(buf, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


PROFILE = {"full_name": "Ada Lovelace", "title": "Engineer", "email": "ada@example.com"}


def _design(**photo) -> ResumeDesign:
    design = ResumeDesign()
    design.layout.photo = ResumePhoto(data_url=_png_data_url(), **photo)
    return design


def _headshots(element) -> list:
    return [d for d in element.iter(qn("wp:inline")) if d.find(qn("wp:docPr")).get("descr") == "Headshot photo"]


def _shape(inline) -> str:
    return inline.find(".//" + qn("a:prstGeom")).get("prst")


def test_photo_rejects_non_image_data_url() -> None:
    with pytest.raises(ValidationError):
        ResumePhoto(data_url="data:text/html;base64,PHNjcmlwdD4=")
    with pytest.raises(ValidationError):
        ProfilePhotoRequest(data_url="https://example.com/me.jpg")
    assert ProfilePhotoRequest(data_url=_png_data_url()).data_url.startswith("data:image/png")


def test_left_header_puts_text_cell_before_photo_cell() -> None:
    doc = Document()
    _render_header(doc, _design(shape="circle"), PROFILE)

    text_cell, photo_cell = doc.tables[0].rows[0].cells
    assert "Ada Lovelace" in text_cell.text
    assert not _headshots(text_cell._tc)
    shots = _headshots(photo_cell._tc)
    assert len(shots) == 1
    assert _shape(shots[0]) == "ellipse"
    assert not photo_cell.text.strip()


def test_center_header_places_photo_above_name() -> None:
    design = _design(shape="rounded")
    design.layout.header_align = "center"
    doc = Document()
    _render_header(doc, design, PROFILE)

    first, second = doc.paragraphs[0], doc.paragraphs[1]
    shots = _headshots(first._p)
    assert len(shots) == 1
    assert _shape(shots[0]) == "roundRect"
    assert second.text == "Ada Lovelace"


@pytest.mark.parametrize("enabled,include", [(False, True), (True, False)])
def test_hidden_or_excluded_photo_renders_plain_header(enabled: bool, include: bool) -> None:
    doc = Document()
    _render_header(doc, _design(enabled=enabled), PROFILE, include_photo=include)
    assert not _headshots(doc.element.body)
    assert doc.paragraphs[0].text == "Ada Lovelace"


def test_band_estimate_grows_for_photo_only_when_included() -> None:
    plain = _header_band_height_pt(ResumeDesign(), PROFILE)
    design = _design(size="lg")
    assert _header_band_height_pt(design, PROFILE, include_photo=False) == plain
    with_photo = _header_band_height_pt(design, PROFILE)
    assert with_photo > plain
    assert with_photo >= PHOTO_SIZE_PT["lg"]
