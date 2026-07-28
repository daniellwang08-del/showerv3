"""Contact hyperlinks in generated resume DOCX headers."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

from app.services.resume_design_compiler import (
    _add_hyperlink_text,
    _clean_url,
    _contact_href,
    _ensure_http_url,
    _render_header,
)
from app.models.resume_design_schemas import ResumeDesign


def test_ensure_http_url_adds_scheme():
    assert _ensure_http_url("linkedin.com/in/chujia-liu") == "https://linkedin.com/in/chujia-liu"
    assert _ensure_http_url("https://www.linkedin.com/in/chujia-liu") == "https://www.linkedin.com/in/chujia-liu"


def test_contact_href_linkedin_and_email():
    assert _contact_href("linkedin", "linkedin.com/in/chujia-liu", "linkedin.com/in/chujia-liu") == (
        "https://linkedin.com/in/chujia-liu"
    )
    assert _contact_href("email", "a@b.com", "a@b.com") == "mailto:a@b.com"
    assert _contact_href("phone", None, "+1 (628) 900-2874") == "tel:+16289002874"


def test_add_hyperlink_text_creates_external_relationship():
    doc = Document()
    p = doc.add_paragraph()
    _add_hyperlink_text(
        p,
        "https://www.linkedin.com/in/chujia-liu",
        "linkedin.com/in/chujia-liu",
        font="Calibri",
        size_pt=10,
        color=RGBColor(0x64, 0x74, 0x8B),
    )
    hyperlinks = list(p._p.iter(qn("w:hyperlink")))
    assert len(hyperlinks) == 1
    r_id = hyperlinks[0].get(qn("r:id"))
    assert r_id
    rel = p.part.rels[r_id]
    assert rel.target_ref == "https://www.linkedin.com/in/chujia-liu"
    assert "linkedin.com/in/chujia-liu" in p.text


def test_render_header_linkedin_is_hyperlink():
    design = ResumeDesign()
    doc = Document()
    profile = {
        "full_name": "Chujia Liu",
        "title": "Staff Software Engineer",
        "email": "liuchujia3@gmail.com",
        "phone": "+1 (628) 900-2874",
        "linkedin": "linkedin.com/in/chujia-liu",
        "github": "",
    }
    _render_header(doc, design, profile)

    found = False
    for p in doc.paragraphs:
        for hl in p._p.iter(qn("w:hyperlink")):
            r_id = hl.get(qn("r:id"))
            if not r_id:
                continue
            target = p.part.rels[r_id].target_ref
            if "linkedin.com/in/chujia-liu" in (target or ""):
                found = True
                assert _clean_url(target) == "linkedin.com/in/chujia-liu" or "linkedin.com/in/chujia-liu" in p.text
    assert found, "Expected a LinkedIn hyperlink in the rendered header"

    # Round-trip through bytes to ensure the relationship survives serialization
    buf = BytesIO()
    doc.save(buf)
    reloaded = Document(BytesIO(buf.getvalue()))
    targets = []
    for p in reloaded.paragraphs:
        for hl in p._p.iter(qn("w:hyperlink")):
            r_id = hl.get(qn("r:id"))
            if r_id:
                targets.append(p.part.rels[r_id].target_ref)
    assert any("linkedin.com/in/chujia-liu" in t for t in targets)
